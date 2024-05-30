import os
import openpyxl
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# 配置部分
folder_path = r'C:\Users\SFZDSB\Desktop\文章5'  # 文章所在的文件夹路径
excel_path = r'C:\Users\SFZDSB\Desktop\domain name.xlsx'  # 存储域名的Excel文件路径
domain_column = 'A'  # 域名所在的Excel列
articles_per_domain = 5  # 每个域名添加到几篇文章中

def add_hyperlink(paragraph, url, text, color="0000FF", underline=True):
    # Create the w:hyperlink tag and add needed values
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    # Create a w:r element and a new w:rPr element
    new_run = OxmlElement('w:r')
    r_pr = OxmlElement('w:rPr')

    # Add color if it is given
    if color:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color)
        r_pr.append(c)

    # Remove or keep underline
    if not underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), "none")
        r_pr.append(u)

    new_run.append(r_pr)
    new_run.text = text
    hyperlink.append(new_run)
    paragraph._element.append(hyperlink)

    return hyperlink

# 读取Excel文件中的域名
wb = openpyxl.load_workbook(excel_path)
sheet = wb.active
domains = [cell.value for cell in sheet[domain_column] if cell.value]

# 获取文件夹中的所有.docx文章文件
articles = [os.path.join(folder_path, f) for f in os.listdir(folder_path) if f.endswith('.docx')]

# 确保域名数量和文章数量匹配
if len(domains) * articles_per_domain != len(articles):
    raise ValueError("文章数量与域名数量不匹配。请检查配置。")

# 为每篇文章添加对应的域名
for i, article_path in enumerate(articles):
    document = Document(article_path)
    domain = domains[i // articles_per_domain]
    paragraph = document.add_paragraph()
    add_hyperlink(paragraph, domain, f"想了解更多文章：{domain}")
    document.save(article_path)

print("操作完成，每篇文章已添加相应的超链接域名。")
